#!/usr/bin/env python3
"""
Editable architecture diagrams as NATIVE PowerPoint shapes (no images):
  slide 1 — ELM–PFLOTRAN weak coupling (vertical, like the physical column)
  slide 2 — current framework (four agents + data layer), hand-editable port

    module load pytorch/2.8.0
    python3 tools/make_diagrams_pptx.py            # -> docs/slides/diagrams_editable.pptx
"""
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn

BLUE = RGBColor(0x2E, 0x75, 0xB6)
INK = RGBColor(0x11, 0x18, 0x27)
MUT = RGBColor(0x47, 0x55, 0x69)
GRAY = RGBColor(0xEF, 0xEF, 0xEF)
GBRD = RGBColor(0x8A, 0x8A, 0x8A)
ORG = RGBColor(0xC2, 0x41, 0x0C)
WHT = RGBColor(0xFF, 0xFF, 0xFF)
W, H = Inches(13.333), Inches(7.5)


def _arrowhead(line):
    ln = line._get_or_add_ln()
    tail = ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"})
    ln.append(tail)


def _dash(line):
    ln = line._get_or_add_ln()
    ln.append(ln.makeelement(qn("a:prstDash"), {"val": "dash"}))


def box(sl, x, y, w, h, fill=WHT, border=BLUE, bw=2.2, round_=True):
    shp = sl.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if round_ else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid(); shp.fill.fore_color.rgb = fill
    shp.line.color.rgb = border; shp.line.width = Pt(bw)
    shp.shadow.inherit = False
    shp.text_frame.word_wrap = True
    return shp


def txt(sl, x, y, w, h, runs, size=11, align=PP_ALIGN.LEFT):
    """runs: list of lines; a line may be (text, {bold,color,size}) or str."""
    tb = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True
    for i, line in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        parts = line if isinstance(line, list) else [line]
        for part in parts:
            t, st = part if isinstance(part, tuple) else (part, {})
            r = p.add_run(); r.text = t
            r.font.size = Pt(st.get("size", size))
            r.font.bold = st.get("bold", False)
            r.font.italic = st.get("italic", False)
            r.font.color.rgb = st.get("color", INK)
    return tb


def fill_box(shp, title, lines, tsize=14, bsize=10.5, tcolor=INK):
    tf = shp.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = title
    r.font.bold = True; r.font.size = Pt(tsize); r.font.color.rgb = tcolor
    for line in lines:
        p = tf.add_paragraph(); p.alignment = PP_ALIGN.LEFT
        parts = line if isinstance(line, list) else [line]
        for part in parts:
            t, st = part if isinstance(part, tuple) else (part, {})
            r = p.add_run(); r.text = t
            r.font.size = Pt(st.get("size", bsize))
            r.font.bold = st.get("bold", False)
            r.font.italic = st.get("italic", False)
            r.font.color.rgb = st.get("color", MUT)


def arrow(sl, x1, y1, x2, y2, color=INK, wpt=1.75, dashed=False, elbow=False):
    c = sl.shapes.add_connector(
        MSO_CONNECTOR.ELBOW if elbow else MSO_CONNECTOR.STRAIGHT,
        Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color; c.line.width = Pt(wpt)
    _arrowhead(c.line)
    if dashed:
        _dash(c.line)
    c.shadow.inherit = False
    return c


# ═════════════════════════ slide 1 — weak coupling ═════════════════════════
def coupling_slide(prs):
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    txt(sl, 0.4, 0.15, 9.6, 0.5,
        [("ELM–PFLOTRAN weak coupling  (per column)", {"bold": True, "size": 24})])
    txt(sl, 0.4, 0.62, 9.0, 0.3,
        [("sequential / offline exchange — each model keeps its own solver, "
          "timestep and code base; files are the interface", {"italic": True,
                                                              "color": MUT})], size=11)

    # atmosphere strip
    atm = box(sl, 3.0, 1.0, 6.4, 0.52, fill=GRAY, border=GBRD, bw=1)
    fill_box(atm, "NLDAS-2 met forcing — 0.125° / 12 km (rain·snow·T·wind·humidity·radiation)",
             [], tsize=11)
    arrow(sl, 6.2, 1.52, 6.2, 1.86)

    # ELM box
    elm = box(sl, 3.0, 1.9, 6.4, 1.55)
    fill_box(elm, "ELM — E3SM Land Model  (surface & root zone, 0–3.8 m + bucket aquifer)",
             [("• snow & canopy · evapotranspiration · runoff partitioning", {}),
              ("• per-column SSURGO soil surface data · NLDAS forcing", {}),
              ("• warm-started at the Fan water table (finidat)", {}),
              [("→ outputs: daily infiltration QINFL · runoff · restart states",
                {"bold": True, "color": INK})]], tsize=13)

    # coupling interface band
    arrow(sl, 5.0, 3.45, 5.0, 4.32, color=BLUE, wpt=2.5)
    txt(sl, 3.15, 3.62, 1.7, 0.6,
        [[("① one-way — LIVE", {"bold": True, "color": BLUE})],
         ("daily QINFL → top", {"color": MUT}),
         ("Neumann flux BC", {"color": MUT})], size=9.5)
    arrow(sl, 7.5, 4.32, 7.5, 3.45, color=ORG, wpt=2.25, dashed=True)
    txt(sl, 7.7, 3.62, 2.2, 0.75,
        [[("② feedback — PLANNED", {"bold": True, "color": ORG})],
         ("water-table depth / capillary", {"color": MUT}),
         ("rise → ELM bottom BC;", {"color": MUT}),
         ("iterate until WTD converges", {"color": MUT})], size=9.5)

    # PFLOTRAN box
    pf = box(sl, 3.0, 4.36, 6.4, 1.75)
    fill_box(pf, "PFLOTRAN — Richards vadose zone  (variably saturated, 1-D to ~50 m)",
             [("• van Genuchten curves from the same SSURGO horizons", {}),
              ("• hydrostatic IC + bottom pinned at the Fan water table", {}),
              ("• ~1 s per column, serial — no batch queue", {}),
              [("→ outputs: recharge AT the water table · lag & attenuation · "
                "saturation profiles", {"bold": True, "color": INK})]], tsize=13)

    # water table strip
    wt = box(sl, 3.0, 6.42, 6.4, 0.52, fill=GRAY, border=GBRD, bw=1)
    fill_box(wt, "water table — Fan 2013 equilibrium prior (ParFlow-CONUS2 upgrade wired)",
             [], tsize=11)
    arrow(sl, 6.2, 6.11, 6.2, 6.42)

    # depth scale (left)
    for y, lab in ((1.9, "0 m"), (3.45, "3.8 m"), (6.11, "~50 m (cap)")):
        txt(sl, 2.1, y - 0.12, 0.85, 0.3, [(lab, {"color": MUT})], size=9,
            align=PP_ALIGN.RIGHT)
    ln = sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(2.95), Inches(1.9),
                                 Inches(2.95), Inches(6.42))
    ln.line.color.rgb = GBRD; ln.line.width = Pt(1)

    # right panel: definition + evidence + legend
    side = box(sl, 10.15, 1.0, 2.95, 5.1, fill=WHT, border=GBRD, bw=1.2)
    fill_box(side, "why WEAK coupling",
             [("• no source-code intrusion — either model can be swapped or "
               "upgraded independently", {}),
              ("• each model runs at its own resolution & timestep", {}),
              ("• exchange = per-column files (flux series, restart states)", {}),
              ("", {}),
              ("evidence (Naches, 18/20 columns):", {"bold": True, "color": INK}),
              ("• infiltration→water-table lag: 0 d (WTD<1 m) → 7–104 d (5–35 m)", {}),
              ("• vadose zone = low-pass filter; cutoff set by WTD", {}),
              ("", {}),
              ("legend:", {"bold": True, "color": INK}),
              ("— solid blue: implemented (one-way)", {"color": BLUE}),
              ("- - dashed orange: planned (two-way)", {"color": ORG})],
             tsize=12, bsize=9.5)

    # shared context footer
    foot = box(sl, 0.4, 7.05, 12.5, 0.38, fill=GRAY, border=GBRD, bw=0.75, round_=False)
    fill_box(foot, "", [("shared per-column context (columns.json): same site, same "
                         "SSURGO soil profile, same Fan-WTD prior — the models see one "
                         "consistent world", {"size": 10, "color": MUT})], tsize=1)


# ═════════════════════ slide 2 — framework (editable port) ══════════════════
def framework_slide(prs):
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    txt(sl, 0.4, 0.1, 12.5, 0.45,
        [("Multi-agent framework — current architecture", {"bold": True, "size": 22})])

    # data layer
    dl = box(sl, 0.3, 0.65, 2.9, 2.6)
    fill_box(dl, "Data Layer",
             [("MCP servers (agent tools):", {"bold": True, "color": INK}),
              ("terrain 3DEP/WBD · geology SSURGO · usgs_water wells+gauges · "
               "snotel SWE · weather · fan_wtd", {}),
              ("Bulk products (build time):", {"bold": True, "color": ORG}),
              ("NLDAS-2 forcing 12 km (default) · Fan WTD warm-start prior · "
               "ParFlow-CONUS2 (wired)", {})], tsize=13, bsize=9.5)

    # reception
    rc = box(sl, 5.7, 0.65, 3.3, 1.5)
    fill_box(rc, "Reception Agent",
             [("agentic LLM tool-user — drives MCP itself", {"italic": True}),
              ("intent + archetype routing · evidence gathering · interactive "
               "clarification · multi-turn context", {})], tsize=13, bsize=9.5)

    # user
    txt(sl, 11.0, 0.65, 2.0, 0.4, [("USER", {"bold": True, "size": 15})])
    arrow(sl, 11.0, 0.95, 9.05, 1.1)
    txt(sl, 9.6, 0.62, 1.5, 0.3, [("question", {"color": MUT})], size=9)
    arrow(sl, 9.05, 1.6, 11.0, 1.15, dashed=True)
    txt(sl, 9.6, 1.55, 1.6, 0.3, [("clarifications", {"color": MUT})], size=9)

    # tool loop
    arrow(sl, 3.25, 1.1, 5.65, 1.1)
    arrow(sl, 5.65, 1.55, 3.25, 1.55)
    txt(sl, 3.6, 0.72, 2.0, 0.3, [("agentic tool loop", {"bold": True})], size=9)

    # planner
    pl = box(sl, 0.3, 3.7, 3.4, 3.1)
    fill_box(pl, "Planner Agent",
             [("capability-aware LLM designer", {"italic": True}),
              ("• strategy: N, strata, validation targets", {}),
              ("• feasibility verdict vs execution limits", {}),
              ("• requires_capabilities backlog", {}),
              ("• altitude only — NO coordinates", {"bold": True, "color": ORG}),
              ("• internal validation loop", {})], tsize=13, bsize=9.5)

    # experiment manager
    em = box(sl, 4.55, 3.7, 4.1, 3.1)
    fill_box(em, "Experiment Manager",
             [("deterministic pipeline — strategy in, results out", {"italic": True}),
              ("ELM per-column · PFLOTRAN 1-D — same interface", {"italic": True}),
              ("1. materialize sampling (DEM·soil·WTD)", {}),
              ("2. build & configure (forcing · warm start)", {}),
              ("3. pre-flight guards", {}),
              ("4. execute (concurrent · cost confirm)", {}),
              ("5. collect (history · restarts)", {})], tsize=13, bsize=9.5)

    # analyzer
    an = box(sl, 9.5, 3.7, 3.5, 3.1)
    fill_box(an, "Analyzer",
             [("numbers → validation → LLM interpretation", {"italic": True}),
              ("a. water budgets · driver×response matrix · honesty flags", {}),
              ("b. observation validation: wells · gauges (NSE/KGE) · SNOTEL", {}),
              ("c. LLM interpreter (Answer/Why/Trust/Next)", {}),
              ("→ auto-generated storyline deck", {"bold": True, "color": INK})],
             tsize=13, bsize=9.5)

    # flows
    arrow(sl, 5.7, 2.15, 2.2, 3.7)
    txt(sl, 3.3, 2.75, 2.3, 0.3, [("framed brief JSON", {"bold": True})], size=9)
    arrow(sl, 3.7, 5.25, 4.55, 5.25, wpt=2.25)
    txt(sl, 3.72, 4.95, 0.9, 0.3, [("strategy", {"bold": True})], size=9)
    arrow(sl, 8.65, 5.25, 9.5, 5.25, wpt=2.25)
    txt(sl, 8.67, 4.95, 0.9, 0.3, [("results", {"bold": True})], size=9)
    arrow(sl, 9.0, 2.15, 11.0, 3.7)
    txt(sl, 9.9, 2.75, 2.3, 0.3, [("analyze existing exp.", {"color": MUT})], size=9)
    arrow(sl, 12.6, 3.7, 12.0, 1.0, dashed=True)
    txt(sl, 12.0, 2.3, 1.3, 0.5, [("final report + deck", {"color": MUT})], size=9)
    arrow(sl, 1.75, 3.25, 5.5, 4.4, dashed=True, color=GBRD, wpt=1.2)
    txt(sl, 1.6, 3.35, 2.6, 0.3, [("DEM·soil·WTD·forcing (bulk)", {"color": MUT})], size=8.5)
    arrow(sl, 3.2, 3.0, 10.5, 3.65, dashed=True, color=GBRD, wpt=1.2)
    txt(sl, 6.6, 3.15, 3.2, 0.3, [("observations for validation", {"color": MUT})], size=8.5)

    # anti-hallucination boundary
    ln = sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(4.15), Inches(3.6),
                                 Inches(4.15), Inches(6.9))
    ln.line.color.rgb = RGBColor(0xDC, 0x26, 0x26); ln.line.width = Pt(1.5)
    _dash(ln.line)
    txt(sl, 2.6, 6.92, 8.5, 0.3,
        [("anti-hallucination boundary — the LLM plans strategy; every coordinate "
          "& config comes from data", {"color": RGBColor(0xDC, 0x26, 0x26),
                                       "bold": True})], size=9)

    # chained experiments loop
    arrow(sl, 8.0, 6.8, 5.2, 6.8, color=ORG, wpt=1.75, dashed=True, elbow=True)
    txt(sl, 5.4, 7.05, 7.4, 0.35,
        [("chained experiments: a run's outputs (ELM infiltration, restarts) feed the "
          "next build — warm starts · ELM→PFLOTRAN coupling", {"color": ORG})], size=9)


def main():
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    coupling_slide(prs)
    framework_slide(prs)
    out = Path(__file__).resolve().parents[1] / "docs" / "slides" / "diagrams_editable.pptx"
    prs.save(str(out))
    print(f"{len(prs.slides._sldIdLst)} slides -> {out}")


if __name__ == "__main__":
    main()
