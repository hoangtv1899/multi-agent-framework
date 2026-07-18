#!/usr/bin/env python3
"""
Assemble the main figures and their captions into a Word manuscript draft.

Style: plain and direct. Short sentences. No em-dashes. Captions state what
is shown and where the numbers come from. Interpretation belongs in the text,
not here.

    python3 tools/make_figure_docx.py   # -> docs/paper/figures_manuscript.docx
"""
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = Path(__file__).resolve().parents[1]
PAPER = ROOT / "docs" / "paper"

# GMD title format: "Name vX.Y: description". NAME is a pending user decision;
# [NAME] is the placeholder until then.
TITLE = ("[NAME] v1.0: a capability-aware multi-agent framework for "
         "trustworthy, observation-validated hydrologic simulation from "
         "natural-language questions (ELM and PFLOTRAN)")

FIGURES = [
    ("fig1_framework.png",
     "Figure 1. Architecture of the framework. The reception agent gathers "
     "evidence through the MCP data servers and produces a framed brief. The "
     "planner produces a sampling strategy and a feasibility verdict. The "
     "planner does not produce coordinates or model configurations. The "
     "experiment manager materializes the strategy against real terrain, "
     "soil, and water-table data, builds and runs the simulations, and "
     "collects the results. The analyzer computes water budgets and driver "
     "relations, validates against observations, and produces an "
     "interpretation. Outputs of one run can feed the next build. This "
     "supports warm starts and the one-way ELM to PFLOTRAN coupling."),

    ("fig2_sampling_design.png",
     "Figure 2. Study area and sampling design for the Naches sub-watershed "
     "(HUC 17030002). Top row: the 20 columns on the DEM, colored by "
     "elevation band, with the watershed boundary in navy and a location "
     "inset; the basin hypsometry with band edges (dashed) and column "
     "elevations (ticks); and the soil configurations sampled from SSURGO. "
     "Bottom row: the Fan (2013) water-table depth at each column; the "
     "number of columns allocated per elevation band; and the NLDAS-2 "
     "annual precipitation for 1995 at each column. All coordinates and "
     "soil values come from the data servers, not from the language model."),

    ("fig3_validation.png",
     "Figure 3. Observation validation for the cold start (top row) and the "
     "Fan warm start (bottom row). The two runs use identical columns, "
     "forcing, and year. (a, d) Water-table depth for the model columns, "
     "the Fan (2013) product at the same locations, and observed USGS "
     "wells. Horizontal lines show medians. (b, e) Modeled water yield and "
     "observed specific discharge at in-domain gauges. No in-domain gauge "
     "returned daily records during the warm-start retrieval, so that panel "
     "is context only. (c, f) Peak snow water equivalent for water year "
     "1995. Bars show the six SNOTEL stations. The shaded band shows the "
     "range across model columns."),

    ("fig4_cold_vs_warm.png",
     "Figure 4. Effect of initialization on recharge and the water table. "
     "The two runs use the same 20 columns, NLDAS-2 forcing, and year 1995. "
     "Only the initial water table differs. (a) Annual recharge per column "
     "for the cold start and the Fan warm start. Negative values indicate "
     "aquifer discharge toward the root zone. (b) Water-table state under "
     "the warm start. Circles show the Fan (2013) prior. Squares show the "
     "initial model state. Crosses show the end-of-year state. The dashed "
     "line shows the uniform cold-start initial state at 8.8 m."),

    ("fig5_lowpass.png",
     "Figure 5. The vadose zone acts as a low-pass filter with a cutoff set "
     "by water-table depth. (a) Change in mean vadose-zone saturation "
     "between the 300 and 30 mm per year recharge scenarios in the "
     "standalone PFLOTRAN sweep. Triangles mark columns whose water table "
     "lies below the 50 m domain cap. (b) Lag between surface infiltration "
     "and delivery at the water table in the coupled ELM to PFLOTRAN run. "
     "Thirteen columns show a coherent response. (c) Attenuation of the "
     "infiltration signal at the water table in the coupled run. "
     "Correlations are computed against log10 water-table depth."),

    ("fig6_agent_eval.png",
     "Figure 6. Planning-stage evaluation of the agent architecture on 20 "
     "pre-registered prompts. All five arms use Claude Opus 4.8 at "
     "temperature 0, and only the architecture differs between arms. The "
     "naive LLM receives the question only. The informed LLM also receives "
     "the same real-data brief the SAGE-Hydro planner gets. The boundary "
     "ablation removes the rule that the planner may only describe strategy "
     "and never write numbers. The limits ablation removes the inventory of "
     "what the model stack can and cannot do. Panels show the fraction of "
     "prompts where an arm emitted coordinates at planning time (a) or an "
     "invalid runnable configuration (b), the fraction of feasibility "
     "verdicts claiming more than the question allows (c), and the measured "
     "total tokens per planning call (d). Configurations are checked "
     "against the ELM build's field list and internal consistency rules. "
     "SAGE-Hydro emits no coordinates and no runnable configuration at "
     "planning time by design. Its single over-claim rates an impossible "
     "2050 projection as partial while explicitly refusing the projection "
     "itself. Provably out-of-basin coordinates were emitted by the naive "
     "arm on 6 of 20 prompts and by no other arm. Conservative errors are "
     "not shown: SAGE-Hydro hedges answerable questions on 11 of 20 "
     "prompts, the ablations on 10, the baselines on 4 to 6. Token counts "
     "are gateway-reported from a declared instrumentation re-run; a "
     "planning call costs on the order of cents at current API prices, and "
     "a full campaign uses about a dozen calls. Per-class verdict "
     "accuracies for both model waves are in the supplement."),
]


def main():
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(11)

    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = t.add_run(TITLE)
    run.bold = True
    run.font.size = Pt(14)

    a = doc.add_paragraph()
    a.alignment = WD_ALIGN_PARAGRAPH.CENTER
    a.add_run("Author list and affiliations to be added.").italic = True

    doc.add_paragraph()
    n = doc.add_paragraph()
    n.add_run("Figures and captions. Draft for internal review.").italic = True

    for fname, caption in FIGURES:
        path = PAPER / fname
        doc.add_page_break()
        if path.exists():
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(str(path), width=Inches(6.5))
        else:
            doc.add_paragraph(f"[missing figure file: {fname}]")
        cap = doc.add_paragraph()
        words = caption.split(". ", 1)
        lead = cap.add_run(words[0] + ". ")
        lead.bold = True
        if len(words) > 1:
            cap.add_run(words[1])
        for run in cap.runs:
            run.font.size = Pt(10)

    out = PAPER / "figures_manuscript.docx"
    doc.save(str(out))
    print(f"✓ {out}")


if __name__ == "__main__":
    main()
