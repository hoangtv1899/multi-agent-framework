#!/usr/bin/env python3
"""
Build a PowerPoint STORYLINE from all executed studies under workflow_outputs/:
for each run —  question → planning figure → planning reasons → execution time
→ result figures → analyzer interpretation (+ observation validation)  — and a
closing cross-study comparison table.

A run is included when it has 04_analysis/hydro_summary.json (i.e. it executed
and was analyzed). Reads only; writes docs/slides/story_<date>.pptx.

    module load pytorch/2.8.0
    python3 tools/make_story_pptx.py [--out <file.pptx>]
"""
import argparse
import json
import re
from datetime import date
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

ROOT = Path(__file__).resolve().parents[1]
W, H = Inches(13.333), Inches(7.5)
INK, MUT, ACC = RGBColor(0x0F, 0x17, 0x2A), RGBColor(0x47, 0x55, 0x69), RGBColor(0x2C, 0x7F, 0xB8)


# ── run discovery + fact extraction ─────────────────────────────────────────
def load(p):
    try:
        return json.loads(Path(p).read_text())
    except Exception:
        return {}


def exec_summary(rd: Path):
    """Parse whatever run log exists for per-column times / total wall time."""
    n_cases = len(load(rd / "cases.json") or load(rd / "phase3_cases.json") or [])
    for logname in ("run.log", "run10_run.log", "phase3_run.log"):
        f = rd / logname
        if not f.exists():
            continue
        txt = f.read_text(errors="ignore")
        cols = re.findall(r"(\S+): rc=(\d+)\s+(\d+)min", txt)
        done = re.search(r"ALL_DONE(?: in (\d+)min)?", txt)
        if cols:
            n_ok = sum(1 for _, rc, _ in cols if rc == "0")
            mins = [int(m) for _, _, m in cols]
            return (f"{n_ok}/{len(cols)} columns succeeded · ~{max(mins)} min each "
                    f"(~15 s clone per column, shared executable)")
        if done:
            t = f" in {done.group(1)} min" if done.group(1) else ""
            return f"{n_cases} columns ran as one batch job (concurrent){t}"
    return f"{n_cases} columns built (no run log found)" if n_cases else None


def forcing_label(rd: Path) -> str:
    """Which met forcing drove this run: explicit marker file, else detect from
    a surviving case's DATM stream, else a dirname heuristic (purged runs)."""
    m = rd / "forcing.txt"
    if m.exists():
        return m.read_text().strip()
    for cf in ("cases_all.json", "cases.json"):
        for c in load(rd / cf) or []:
            st = Path(c) / "run" / "datm.streams.txt.CLM_QIAN.Precip"
            if st.exists():
                return ("NLDAS-2 (0.125°, ~12 km)" if "NLDAS2" in st.read_text()
                        else "CLM_QIAN (Qian 2006, T62 ≈ 1.9°)")
    return ("NLDAS-2 (0.125°, ~12 km)" if "nldas" in rd.name or "warmstart" in rd.name
            else "CLM_QIAN (Qian 2006, T62 ≈ 1.9°)")


def gather(rd: Path):
    """Everything the slides need for one run."""
    hs = load(rd / "04_analysis" / "hydro_summary.json")
    if not hs:
        return None
    brief = load(rd / "reception_brief.json")
    plan = load(rd / "plan.json")
    dom = brief.get("domain") or {}
    sweep = load(rd / "soilsweep_plan.json")

    if brief:
        question = brief.get("user_request", "?")
        name = dom.get("name") or rd.name
        rp = load(rd / "run_plan.json")
        if any(cc.get("FINIDAT") for cc in (rp.get("CONDITIONS_COUPLERS") or [])):
            name += " — Fan warm start"
        subtitle = (f"HUC {dom.get('huc', '?')} · {dom.get('area_km2', '?')} km² · "
                    f"elevation {dom.get('elevation_range_m', {}).get('min', '?')}–"
                    f"{dom.get('elevation_range_m', {}).get('max', '?')} m")
    elif sweep:
        cc = sweep.get("CONDITIONS_COUPLERS", [])
        clays = [c["EXPERIMENT"] for c in cc]
        question = ("Controlled experiment: how does soil texture (clay content) "
                    "control the runoff/recharge partitioning?")
        name = "Controlled soil sweep"
        subtitle = (f"{len(cc)} uniform-soil treatments ({', '.join(clays)}) at one "
                    f"site — forcing, vegetation and location held identical")
    else:
        question, name, subtitle = "?", rd.name, ""

    # planning reasons
    reasons = []
    sd = plan.get("scientific_decomposition") or {}
    for g in (sd.get("goals") or [])[:2]:
        reasons.append(("goal", g))
    ss = plan.get("sampling_strategy") or {}
    if ss:
        reasons.append(("sampling", f"N={ss.get('n_exploratory')} — {ss.get('approach', '')}"))
        if ss.get("n_justification"):
            reasons.append(("why N", ss["n_justification"]))
    fe = plan.get("feasibility") or {}
    if fe:
        reasons.append(("feasibility", str(fe.get("verdict", ""))))
        for x in (fe.get("not_answerable") or [])[:2]:
            if isinstance(x, dict):
                reasons.append(("not answerable", f"{x.get('aspect')} — needs {x.get('needs')}"))
    if sweep and not plan:
        reasons.append(("design", "conceptual archetype — vary ONE factor (clay 5→55%), "
                                  "hold everything else fixed; the cleanest isolation of soil control"))

    # analyzer interpretation — prefer the LLM's (interpret_run.py), else the
    # deterministic notes. Render markdown cleanly: **Section** -> bold header
    # tuple, bullets -> plain lines (no stray ** / # / - markers).
    interp = []
    llm_md = rd / "04_analysis" / "interpretation.md"
    if llm_md.exists():
        for line in llm_md.read_text().splitlines():
            raw = line.strip()
            if not raw:
                continue
            t = raw.replace("*", "").lstrip("#-• ").strip()
            if not t:
                continue
            if t.rstrip(":") in ("Answer", "Why", "Trust", "Next"):
                interp.append((t.rstrip(":"), ""))        # bold section header
            else:
                interp.append("•  " + t)
        interp = interp[:18]
    ssum = hs.get("spatial_summary") or {}
    sa = hs.get("soil_attribution") or {}
    if not interp:
        for n in ssum.get("interpretation") or []:
            interp.append(n)
        if sa:
            interp.append(f"soil control (forcing held at {sa.get('forcing_held_mm_yr')} mm/yr, "
                          f"{sa.get('n_columns')} columns): strongest predictor = "
                          f"{sa.get('strongest_predictor')}")
    # evaluation vs observations (its own slide)
    val = load(rd / "04_analysis" / "validation.json")
    evaluation = None
    if val:
        inv = val.get("observation_inventory") or {}
        lines = [("in-domain obs",
                  f"SNOTEL {inv.get('snotel', {}).get('n', '?')} · "
                  f"stream gages {inv.get('streamgages', {}).get('n', '?')} · "
                  f"GW wells {inv.get('gwwells', {}).get('n', '?')} "
                  f"({inv.get('gwwells', {}).get('with_records', '?')}/"
                  f"{inv.get('gwwells', {}).get('sampled', '?')} sampled have records)")]
        for t in val.get("targets") or []:
            tag = "✓ COMPARED" if t.get("status") == "compared" else "· context-only"
            body = t.get("result", "")
            if t.get("needs"):
                body += f"  — needs: {t['needs']}"
            lines.append((f"{tag} — {t.get('variable')}", body))
        fig = rd / "04_analysis" / "validation.png"
        evaluation = {"lines": lines, "fig": fig if fig.exists() else None}

    figs = [rd / f for f in ("sampling_design.png",) if (rd / f).exists()]
    rfigs = [rd / "04_analysis" / f
             for f in ("elevation_gradient.png", "water_budget.png",
                       "driver_response.png", "wtd_columns.png", "cold_vs_warm.png",
                       "soil_control.png", "debug_timeseries.png")
             if (rd / "04_analysis" / f).exists()]

    ok = [r for r in hs.get("experiments", []) if r.get("status") == "ok"]
    rech = [r["metrics"].get("annual_recharge_mm_yr") for r in ok
            if r["metrics"].get("annual_recharge_mm_yr") is not None]

    # per-column appendix rows: characteristics (columns.json) + run results
    chars = {}
    cj = load(rd / "columns.json")
    for c in (cj.get("columns", cj) if isinstance(cj, (dict, list)) else []) or []:
        if isinstance(c, dict):
            chars[c.get("id")] = c
    def fmt(x, d=1):
        return f"{x:.{d}f}" if isinstance(x, (int, float)) else "—"
    rows = []
    for r in sorted(ok, key=lambda r: (r.get("elevation_m") or 0)):
        ch = chars.get(r["case_name"], {})
        so = r.get("soil") or {}
        m = r["metrics"]
        rows.append([r["case_name"], fmt(r.get("elevation_m"), 0),
                     str(so.get("texture_top") or ch.get("soil_top_texture") or "—"),
                     fmt(so.get("clay_max_pct"), 0), fmt(so.get("ksat_min_ums"), 1),
                     fmt(ch.get("fan_wtd_m"), 1), fmt(m.get("precip_mm_yr"), 0),
                     fmt(m.get("annual_recharge_mm_yr"), 1), fmt(m.get("annual_runoff_mm_yr"), 1),
                     fmt(m.get("recharge_fraction"), 2), fmt(m.get("water_table_depth_m"), 2)])
    return {"dir": rd, "name": name, "question": question, "subtitle": subtitle,
            "reasons": reasons, "execution": exec_summary(rd), "interp": interp,
            "evaluation": evaluation, "column_rows": rows,
            "forcing": forcing_label(rd),
            "driver_matrix": hs.get("driver_matrix"),
            "plan_figs": figs, "result_figs": rfigs,
            "n_cols": len(ok),
            # single-location runs have no spatial summary — fall back to the
            # held-forcing value from the soil attribution
            "precip_bins": ((ssum.get("forcing") or {}).get("precip_mm_yr_distinct")
                            or ([sa["forcing_held_mm_yr"]] if sa.get("forcing_held_mm_yr") else None)),
            "rech_range": (round(min(rech), 1), round(max(rech), 1)) if rech else None,
            "predictor": (sa or {}).get("strongest_predictor")}


# ── pptx helpers ─────────────────────────────────────────────────────────────
def blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def text(slide, x, y, w, h, runs, size=14, color=INK, bold=False, align=None):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    first = True
    for r in (runs if isinstance(runs, list) else [runs]):
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        if isinstance(r, tuple):                      # (label, body); empty body = header
            lab, body = r
            ra = p.add_run(); ra.text = f"{lab}:  " if body else str(lab)
            ra.font.bold = True; ra.font.size = Pt(size); ra.font.color.rgb = ACC
            if body:
                rb = p.add_run(); rb.text = str(body)
                rb.font.size = Pt(size); rb.font.color.rgb = color
        else:
            ra = p.add_run(); ra.text = str(r)
            ra.font.size = Pt(size); ra.font.bold = bold; ra.font.color.rgb = color
        p.space_after = Pt(6)
    return tb


def header(slide, title, sub=None):
    text(slide, Inches(.5), Inches(.25), W - Inches(1), Inches(.7), title, size=28, bold=True)
    if sub:
        text(slide, Inches(.5), Inches(.95), W - Inches(1), Inches(.4), sub, size=13, color=MUT)


def picture(slide, path, top, max_w=Inches(12.3), max_h=None):
    pic = slide.shapes.add_picture(str(path), 0, top, width=max_w)
    if max_h and pic.height > max_h:                  # rescale to height cap
        scale = max_h / pic.height
        pic.width = int(pic.width * scale); pic.height = int(pic.height * scale)
    pic.left = int((W - pic.width) / 2)
    return pic


# ── slide builders ───────────────────────────────────────────────────────────
def gather_pflotran(rd: Path):
    st = load(rd / "pflotran_study.json")
    if not st:
        return None
    return {"type": "pflotran", "dir": rd, "name": st["name"],
            "question": st["question"], "model": st.get("model", ""),
            "execution": st.get("execution", ""), "interp": st.get("interpretation", []),
            "rates": st.get("scenarios_mm_yr", []), "rows": st.get("columns", []),
            "figs": [rd / f for f in ("sampling_design.png",) if (rd / f).exists()],
            "rfigs": [p for p in (rd / "r100" / "pflotran_profiles.png",
                                  rd / "pflotran_scenarios.png") if p.exists()]}


def pflotran_slides(prs, s, idx):
    sl = blank(prs)
    text(sl, Inches(.8), Inches(2.1), W - Inches(1.6), Inches(1.0),
         f"Case {idx} — {s['name']}", size=30, bold=True)
    text(sl, Inches(.8), Inches(3.3), W - Inches(1.6), Inches(2.4),
         [("question", s["question"]), ("model", s["model"]),
          ("scenarios", f"steady uniform recharge {'/'.join(f'{r:.0f}' for r in s['rates'])} mm/yr"
                        " — a controlled experiment: only the subsurface varies"),
          ("execution", s["execution"])], size=15)
    for f in s["figs"]:
        sl = blank(prs)
        header(sl, f"{s['name']} — sampling design (same columns as the ELM studies)")
        sl.shapes.add_picture(str(f), Inches(1.2), Inches(1.25), height=Inches(5.7))
    for cap, f in zip(("profiles & water tables (100 mm/yr)", "scenario sweep — the dynamics"),
                      s["rfigs"]):
        sl = blank(prs)
        header(sl, f"{s['name']} — {cap}")
        sl.shapes.add_picture(str(f), Inches(1.0), Inches(1.35), width=Inches(11.3))
    sl = blank(prs)
    header(sl, f"{s['name']} — interpretation")
    text(sl, Inches(.6), Inches(1.4), W - Inches(1.2), Inches(5.6),
         ["•  " + t for t in s["interp"]], size=14)


def pflotran_appendix(prs, s, idx):
    rows = s.get("rows") or []
    if not rows:
        return
    sl = blank(prs)
    header(sl, f"Appendix P{idx} — {s['name']}: columns & scenario wetness",
           "vadose-zone mean saturation per recharge rate · WTD from the hydrostatic prior")
    rates = s["rates"]
    hdr = ["column", "elev m", "Fan WTD m", "domain m", "WTD end m"] +           [f"sat @{r:.0f}" for r in rates]
    tbl = sl.shapes.add_table(len(rows) + 1, len(hdr), Inches(.7), Inches(1.5),
                              Inches(12), Inches(.26) * (len(rows) + 1)).table
    def fmt(x, d=2):
        return f"{x:.{d}f}" if isinstance(x, (int, float)) else ">cap"
    for c, h in enumerate(hdr):
        cell = tbl.cell(0, c); cell.text = h
        cell.text_frame.paragraphs[0].font.size = Pt(10)
        cell.text_frame.paragraphs[0].font.bold = True
    for r, row in enumerate(rows, 1):
        vals = [row["id"], fmt(row["elevation_m"], 0), fmt(row["fan_wtd_m"], 1),
                fmt(row["depth_m"], 1), fmt(row["wtd_final_m"])] +                [fmt(row.get(f"sat_r{ra:.0f}")) for ra in rates]
        for c, v in enumerate(vals):
            cell = tbl.cell(r, c); cell.text = str(v)
            cell.text_frame.paragraphs[0].font.size = Pt(9)


def study_slides(prs, s, idx):
    # 1 · question + planning reasons
    sl = blank(prs)
    header(sl, f"Case {idx} — {s['name']}", s["subtitle"])
    text(sl, Inches(.5), Inches(1.5), W - Inches(1), Inches(1.1),
         [f"“{s['question']}”"], size=18, bold=True, color=ACC)
    text(sl, Inches(.5), Inches(2.7), W - Inches(1), Inches(4.3),
         [("planner", "")] and s["reasons"] or ["(no plan.json — design-only artifacts absent)"],
         size=14)

    # 2 · planning figure (+ panel definitions)
    for f in s["plan_figs"]:
        sl = blank(prs)
        header(sl, f"{s['name']} — sampling design (Tier-2 expander)")
        picture(sl, f, Inches(1.1), max_h=Inches(5.3))
        text(sl, Inches(.5), Inches(6.5), W - Inches(1), Inches(.95),
             ["Panels — top-left: chosen columns on the DEM, coloured by elevation band, "
              "navy = watershed boundary.  top-right: histogram of the basin's DEM sample "
              "(grey = how much terrain sits at each elevation); dashed lines = the equal-"
              "interval band edges; coloured ticks = the chosen columns' elevations.  "
              "bottom-left: Fan (2013) equilibrium water-table depth at each column.  "
              "bottom-right: columns allocated per band — proportional to the band's share "
              "of the basin's DEM points (annotated n / points), minimum 1 per band."],
             size=10.5, color=MUT)

    # 3 · execution + results figures (+ variable definitions)
    sl = blank(prs)
    header(sl, f"{s['name']} — execution & results")
    if s["execution"]:
        text(sl, Inches(.5), Inches(1.05), W - Inches(1), Inches(.45),
             [("execution", s["execution"])], size=13)
    top = Inches(1.6)
    for f in s["result_figs"][:2]:
        pic = picture(sl, f, top, max_h=Inches(2.6))
        top = pic.top + pic.height + Inches(.1)
    text(sl, Inches(.5), Inches(7.0), W - Inches(1), Inches(.45),
         ["Definitions — recharge = ELM QCHARGE (flux from the soil column to the water "
          "table); runoff = QOVER (surface runoff); recharge fraction = QCHARGE/(QCHARGE"
          "+QOVER); WTD = ZWT (the column's own water-table depth); forcing = "
          f"{s['forcing']} precip at the column."], size=9.5, color=MUT)
    for f in s["result_figs"][2:]:
        sl = blank(prs)
        header(sl, f"{s['name']} — results (continued)")
        picture(sl, f, Inches(1.3), max_h=Inches(5.9))

    # 4 · driver x response matrix (the relationship structure)
    dm = s.get("driver_matrix")
    if dm and dm.get("pearson_r"):
        sl = blank(prs)
        header(sl, f"{s['name']} — driver × response (Pearson r)",
               f"all {dm['n_columns']} columns · "
               + (dm.get("note") or "")[:90])
        drivers = ["elevation_m", "precip_mm_yr", "clay_max_pct", "ksat_min_ums"]
        dlabel = {"elevation_m": "elevation", "precip_mm_yr": "precip",
                  "clay_max_pct": "clay", "ksat_min_ums": "Ksat"}
        rows = [["response"] + [dlabel[d] for d in drivers]]
        for resp, row in dm["pearson_r"].items():
            rows.append([resp] + [("—" if row.get(d) is None else f"{row[d]:+.2f}")
                                  for d in drivers])
        tbl = sl.shapes.add_table(len(rows), 5, Inches(1.2), Inches(1.9),
                                  Inches(10.9), Inches(.55) * len(rows)).table
        for r, row in enumerate(rows):
            for c, v in enumerate(row):
                cell = tbl.cell(r, c)
                cell.text = v
                pr = cell.text_frame.paragraphs[0]
                pr.font.size = Pt(15); pr.font.bold = (r == 0 or c == 0)

    # 5 · analyzer interpretation
    sl = blank(prs)
    header(sl, f"{s['name']} — analyzer interpretation")
    text(sl, Inches(.5), Inches(1.3), W - Inches(1), Inches(5.8),
         s["interp"] or ["(no interpretation recorded)"], size=13)

    # 5 · evaluation vs observations
    if s.get("evaluation"):
        ev = s["evaluation"]
        sl = blank(prs)
        header(sl, f"{s['name']} — evaluation vs observations",
               "what could be compared honestly, and what still blocks rigorous validation")
        text(sl, Inches(.5), Inches(1.25), W - Inches(1), Inches(2.4), ev["lines"], size=12)
        if ev["fig"]:
            picture(sl, ev["fig"], Inches(3.7), max_h=Inches(3.3))
        text(sl, Inches(.5), Inches(7.05), W - Inches(1), Inches(.4),
             ["Metrics — NSE: 1=perfect, 0=no better than the observed mean, <0 worse than "
              "the mean.  KGE: 1=perfect; components r (timing correlation), α (variability "
              "ratio σm/σo), β (volume ratio μm/μo).  Specific discharge = gauge flow ÷ "
              "drainage area, making a gauge comparable to 1-D columns (mm/yr or mm/day)."],
             size=9.5, color=MUT)


def comparison_slide(prs, studies):
    sl = blank(prs)
    header(sl, "Cross-study synthesis", "same framework, three experiments — the control on "
           "recharge/runoff partitioning depends on climate")
    rows = [["study", "columns", "forcing (mm/yr)", "recharge (mm/yr)", "strongest control"]]
    for s in studies:
        rows.append([s["name"], str(s["n_cols"]),
                     ", ".join(str(b) for b in (s["precip_bins"] or [])) or "—",
                     f"{s['rech_range'][0]} → {s['rech_range'][1]}" if s["rech_range"] else "—",
                     str(s["predictor"] or "—")])
    tbl = sl.shapes.add_table(len(rows), 5, Inches(.5), Inches(1.6),
                              W - Inches(1), Inches(.5) * len(rows)).table
    widths = [3.2, 1.2, 2.0, 2.2, 3.7]
    for i, w in enumerate(widths):
        tbl.columns[i].width = Inches(w)
    for r, row in enumerate(rows):
        for c, v in enumerate(row):
            cell = tbl.cell(r, c)
            cell.text = v
            para = cell.text_frame.paragraphs[0]
            para.font.size = Pt(13 if r else 13)
            para.font.bold = (r == 0)
    text(sl, Inches(.5), Inches(1.6) + Inches(.5) * len(rows) + Inches(.4),
         W - Inches(1), Inches(2.2),
         [("wet basin", "recharge is CLAY-limited — water is ample; the impeding layer decides"),
          ("dry basin", "recharge is Ksat-limited (threshold) — only the most permeable soil recharges"),
          ("controlled sweep", "confirms causality: clay 5→55% collapses recharge 755→9 mm/yr "
                               "with everything else fixed")], size=14)


APPENDIX_HDR = ["column", "elev m", "top soil", "clay%", "Ksat µm/s", "Fan WTD m",
                "precip", "recharge", "runoff", "rech.frac", "ZWT m"]


def appendix_slide(prs, s, idx):
    """One appendix slide per study: every sampling point's characteristics
    (location/soil/Fan WTD) next to its simulated water balance."""
    if not s.get("column_rows"):
        return
    sl = blank(prs)
    header(sl, f"Appendix A{idx} — {s['name']}: columns & results",
           f"characteristics from the real data (SSURGO soil, Fan 2013 WTD, "
           f"{s['forcing']} forcing) · results are annual means (mm/yr)")
    rows = [APPENDIX_HDR] + s["column_rows"]
    tbl = sl.shapes.add_table(len(rows), len(APPENDIX_HDR), Inches(.4), Inches(1.45),
                              W - Inches(.8), Inches(.42) * len(rows)).table
    for r, row in enumerate(rows):
        for c, v in enumerate(row):
            cell = tbl.cell(r, c)
            cell.text = str(v)
            para = cell.text_frame.paragraphs[0]
            para.font.size = Pt(11)
            para.font.bold = (r == 0)


def main():
    ap = argparse.ArgumentParser(description="Build the storyline PPTX from executed runs")
    ap.add_argument("--out", default=None)
    args = ap.parse_args()

    run_dirs = sorted(
        d for d in (ROOT / "workflow_outputs").iterdir()
        if (d / "04_analysis" / "hydro_summary.json").exists()
        # a framework case study has a reception brief (spatial) or a sweep plan
        # (controlled) — this excludes legacy pre-framework run dirs
        and ((d / "reception_brief.json").exists() or (d / "soilsweep_plan.json").exists()))
    studies = [s for s in (gather(rd) for rd in run_dirs) if s]
    pf_dirs = sorted(d for d in (ROOT / "workflow_outputs").iterdir()
                     if (d / "pflotran_study.json").exists())
    pf_studies = [s for s in (gather_pflotran(rd) for rd in pf_dirs) if s]
    if not studies:
        raise SystemExit("no executed+analyzed runs found under workflow_outputs/")

    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H

    sl = blank(prs)
    text(sl, Inches(.8), Inches(2.3), W - Inches(1.6), Inches(1.4),
         "From a question to a validated simulation campaign", size=34, bold=True)
    text(sl, Inches(.8), Inches(3.8), W - Inches(1.6), Inches(1.6),
         [f"Multi-agent LLM framework on NERSC Perlmutter — {len(studies)} executed case studies",
          f"question → plan → sampling → per-column ELM runs → analysis → observation validation",
          f"generated {date.today().isoformat()} from workflow_outputs/"], size=16, color=MUT)

    for i, s in enumerate(studies, 1):
        study_slides(prs, s, i)
    for j, s in enumerate(pf_studies, len(studies) + 1):
        pflotran_slides(prs, s, j)
    comparison_slide(prs, studies)
    for i, s in enumerate(studies, 1):          # appendix: per-study column tables
        appendix_slide(prs, s, i)
    for j, s in enumerate(pf_studies, 1):
        pflotran_appendix(prs, s, j)

    out = Path(args.out) if args.out else ROOT / "docs" / "slides" / f"story_{date.today():%Y%m%d}.pptx"
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out))
    print(f"{len(prs.slides)} slides · {len(studies) + len(pf_studies)} studies -> {out}")
    for s in studies + pf_studies:
        print(f"  • {s['name']:<24} ({s['dir'].name})")


if __name__ == "__main__":
    main()
